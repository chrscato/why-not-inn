"""Parse CMS IDR Public Use Files into SQLite.

Walks ``data/puf/`` and ingests every ``.xlsx`` (and ``.zip`` containing
``.xlsx`` or per-tab ``.csv`` files). Three tabs map to two tables:

* ``OON Emergency and Non-Emergency`` -> ``idr_disputes`` (source_tab=oon)
* ``OON Air Ambulance``               -> ``idr_disputes`` (source_tab=air_ambulance)
* ``QPA and Offers``                  -> ``idr_offers``

Column headers vary across quarterly releases (e.g. ``% of QPA`` vs
``Percent of QPA``, ``Initiating Party`` only from late 2023). We use
fuzzy matching: normalize headers, then look up by alias set. Missing
columns become NULL.

Usage:
    python scripts/parse_puf.py                 # ingest everything new
    python scripts/parse_puf.py --force         # reload even if logged
    python scripts/parse_puf.py --file path.xlsx
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sqlite3
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Iterable

import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DB_PATH = PROJECT_ROOT / "db" / "whynotinn.db"
PUF_DIR = PROJECT_ROOT / "data" / "puf"
SCHEMA_PATH = PROJECT_ROOT / "db" / "schema.sql"

CHUNK_SIZE = 5_000


# ---------------------------------------------------------------------------
# Header normalization & alias maps
# ---------------------------------------------------------------------------

def _norm(s: str) -> str:
    """Lowercase, strip, collapse whitespace, drop punctuation."""
    if s is None:
        return ""
    s = str(s).lower()
    s = s.replace("%", " percent ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


# Each entry: db_column -> list of accepted normalized header forms.
# We try aliases in order; first match wins.
TAB_DISPUTES_ALIASES: dict[str, list[str]] = {
    "dispute_number":               [_norm("Dispute Number")],
    "dli_number":                   [_norm("DLI Number")],
    "type_of_dispute":              [_norm("Type of Dispute")],
    "initiating_party":             [_norm("Initiating Party")],
    "default_decision":             [_norm("Default Decision")],
    "payment_determination_outcome":[_norm("Payment Determination Outcome")],
    "length_of_time_days":          [_norm("Length of Time to Make Determination")],
    "idre_compensation":            [_norm("IDRE Compensation")],
    "provider_name":                [_norm("Provider/Facility Name")],
    "provider_group_name":          [_norm("Provider/Facility Group Name")],
    "provider_npi":                 [_norm("Provider/Facility NPI Number")],
    "provider_email_domain":        [_norm("Provider Email Domain")],
    "provider_specialty":           [_norm("Practice/Facility Specialty or Type")],
    "practice_size":                [_norm("Practice/Facility Size")],
    "health_plan_name":             [_norm("Health Plan/Issuer Name")],
    "health_plan_email_domain":     [_norm("Health Plan/Issuer Email Domain")],
    "health_plan_type":             [_norm("Health Plan Type")],
    "service_code":                 [_norm("Service Code")],
    "type_of_service_code":         [_norm("Type of Service Code")],
    "item_description":             [_norm("Item or Service Description")],
    "location_of_service":          [_norm("Location of Service")],
    "place_of_service_code":        [_norm("Place of Service Code")],
    "dispute_line_item_type":       [_norm("Dispute Line Item Type")],
    "offer_selected":               [_norm("Offer Selected from Provider or Issuer")],
    "provider_offer_pct_qpa": [
        _norm("Provider/Facility Offer as % of QPA"),
        _norm("Provider/Facility Offer as Percent of QPA"),
    ],
    "issuer_offer_pct_qpa": [
        _norm("Health Plan/Issuer Offer as % of QPA"),
        _norm("Health Plan/Issuer Offer as Percent of QPA"),
    ],
    "prevailing_offer_pct_qpa": [
        _norm("Prevailing Party Offer as % of QPA"),
        _norm("Prevailing Party Offer as Percent of QPA"),
        _norm("Prevailing Offer as % of QPA"),
        _norm("Prevailing Offer as Percent of QPA"),
    ],
    "provider_offer_pct_median":    [_norm("Provider/Facility Offer as Percent of Median Provider/Facility Offer Amount")],
    "issuer_offer_pct_median":      [_norm("Health Plan/Issuer Offer as Percent of Median Health Plan/Issuer Offer Amount")],
    "prevailing_offer_pct_median":  [_norm("Prevailing Offer as Percent of Median Prevailing Offer Amount")],
    "qpa_pct_median":               [_norm("QPA as Percent of Median QPA")],
    "air_ambulance_vehicle_type": [
        _norm("Air Ambulance Vehicle Type"),
        _norm("Ambulance Vehicle Type"),
    ],
    "air_ambulance_clinical_capacity": [
        _norm("Air Ambulance Vehicle Clinical Capacity Level"),
        _norm("Ambulance Vehicle Clinical Capacity Level"),
    ],
    "air_ambulance_pickup_location": [
        _norm("Air Ambulance Pick-up Location"),
        _norm("Air Ambulance Pickup Location"),
    ],
}

TAB_OFFERS_ALIASES: dict[str, list[str]] = {
    "service_code":                  [_norm("Service Code")],
    "type_of_service_code":          [_norm("Type of Service Code")],
    "geographic_region":             [_norm("Geographical Region"), _norm("Geographic Region")],
    "place_of_service_code":         [_norm("Place of Service Code")],
    "dispute_line_item_type":        [_norm("Dispute Line Item Type")],
    "initiating_party":              [_norm("Initiating Party")],
    "default_decision":              [_norm("Default Decision")],
    "offer_selected":                [_norm("Offer Selected from Provider or Issuer")],
    "qpa":                           [_norm("QPA")],
    "provider_offer":                [_norm("Provider/Facility Offer")],
    "issuer_offer":                  [_norm("Health Plan/Issuer Offer")],
    "prevailing_offer":              [_norm("Prevailing Offer")],
    "air_ambulance_pickup_location": [_norm("Air Ambulance Pick-up Location"), _norm("Air Ambulance Pickup Location")],
}


def build_column_resolver(headers: Iterable[str], alias_map: dict[str, list[str]]) -> dict[str, str | None]:
    """Map db_column -> source header that exists in the dataframe (or None)."""
    headers = list(headers)
    norm_to_orig = {_norm(h): h for h in headers}
    resolved: dict[str, str | None] = {}
    for db_col, aliases in alias_map.items():
        match = None
        for alias in aliases:
            if alias in norm_to_orig:
                match = norm_to_orig[alias]
                break
        resolved[db_col] = match
    return resolved


# ---------------------------------------------------------------------------
# Quarter inference & file discovery
# ---------------------------------------------------------------------------

QUARTER_RE = re.compile(r"(20\d{2})[\s\-_]*q?(\d)(?:[\s\-_]*q?(\d))?", re.IGNORECASE)


def infer_quarter(name: str) -> str:
    """Pull a quarter tag from a filename. Examples:
    ``2023-q1-federal-idr-puf_0.xlsx``                    -> ``2023-Q1``
    ``Federal-IDR-PUF-for-2023-Q2.xlsx``                  -> ``2023-Q2``
    ``federal-idr-puf-for-2024-q3-as-of-may-28-2025.xlsx``-> ``2024-Q3``
    """
    base = Path(name).stem
    # Prefer the FIRST year-quarter token (e.g. ``2024-q3-as-of-2025...``).
    m = QUARTER_RE.search(base)
    if not m:
        return base  # fallback
    year, q1, q2 = m.group(1), m.group(2), m.group(3)
    if q2 and q2 != q1:
        return f"{year}-Q{q1}Q{q2}"
    return f"{year}-Q{q1}"


def classify_sheet(sheet_name: str) -> str | None:
    s = sheet_name.strip().lower()
    if s == "contents":
        return None
    if "air ambulance" in s:
        return "air_ambulance"
    if "qpa" in s and "offer" in s:
        return "offers"
    if "oon" in s or "emergency" in s:
        return "oon"
    return None


# ---------------------------------------------------------------------------
# Insurer name normalization
# ---------------------------------------------------------------------------

INSURER_NORMALIZERS = [
    # (regex, canonical name) — applied in order
    (re.compile(r"united\s*health(care)?", re.I), "UnitedHealthcare"),
    (re.compile(r"\buhc\b", re.I),               "UnitedHealthcare"),
    (re.compile(r"\banthem\b", re.I),            "Anthem"),
    (re.compile(r"\belevance\b", re.I),          "Anthem"),
    (re.compile(r"blue\s*cross.*blue\s*shield|bcbs", re.I), "Blue Cross Blue Shield"),
    (re.compile(r"\baetna\b", re.I),             "Aetna"),
    (re.compile(r"\bcigna\b", re.I),             "Cigna"),
    (re.compile(r"\bhumana\b", re.I),            "Humana"),
    (re.compile(r"kaiser", re.I),                "Kaiser Permanente"),
    (re.compile(r"\bcentene\b", re.I),           "Centene"),
    (re.compile(r"molina", re.I),                "Molina Healthcare"),
    (re.compile(r"highmark", re.I),              "Highmark"),
    (re.compile(r"medica", re.I),                "Medica"),
    (re.compile(r"oscar", re.I),                 "Oscar Health"),
]


def normalize_insurer(name) -> str | None:
    if name is None:
        return None
    s = str(name).strip()
    if not s:
        return None
    for pattern, canonical in INSURER_NORMALIZERS:
        if pattern.search(s):
            return canonical
    # generic cleanup: collapse whitespace, title-case obvious all-caps strings
    s = re.sub(r"\s+", " ", s)
    if s.isupper() and len(s) > 4:
        s = s.title()
    return s


# ---------------------------------------------------------------------------
# Coercion helpers
# ---------------------------------------------------------------------------

_PCT_PATTERN = re.compile(r"-?\d+(?:\.\d+)?")


def to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if pd.isna(v):
            return None
        return float(v)
    s = str(v).strip()
    if not s or s.lower() in {"na", "n/a", "none", "null", "-", "redacted"}:
        return None
    s = s.replace(",", "").replace("$", "").replace("%", "")
    try:
        return float(s)
    except ValueError:
        m = _PCT_PATTERN.search(s)
        return float(m.group(0)) if m else None


def to_int(v):
    f = to_float(v)
    return int(f) if f is not None else None


def to_str(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    s = str(v).strip()
    return s if s else None


# ---------------------------------------------------------------------------
# Row builders
# ---------------------------------------------------------------------------

DISPUTES_COLUMNS = [
    "dispute_number","quarter","source_tab","type_of_dispute","initiating_party",
    "default_decision","payment_determination_outcome","length_of_time_days",
    "idre_compensation","provider_name","provider_group_name","provider_npi",
    "provider_email_domain","provider_specialty","practice_size",
    "health_plan_name","health_plan_name_normalized","health_plan_email_domain",
    "health_plan_type","dli_number","service_code","type_of_service_code",
    "item_description","location_of_service","place_of_service_code",
    "dispute_line_item_type","offer_selected","provider_offer_pct_qpa",
    "issuer_offer_pct_qpa","prevailing_offer_pct_qpa","provider_offer_pct_median",
    "issuer_offer_pct_median","prevailing_offer_pct_median","qpa_pct_median",
    "air_ambulance_vehicle_type","air_ambulance_clinical_capacity",
    "air_ambulance_pickup_location",
]

OFFERS_COLUMNS = [
    "quarter","service_code","type_of_service_code","geographic_region",
    "place_of_service_code","dispute_line_item_type","initiating_party",
    "default_decision","offer_selected","qpa","provider_offer","issuer_offer",
    "prevailing_offer","air_ambulance_pickup_location","prevailing_pct_qpa",
    "qpa_suspect",
]

NUMERIC_DISPUTE_FIELDS = {
    "length_of_time_days":      to_int,
    "idre_compensation":        to_float,
    "provider_offer_pct_qpa":   to_float,
    "issuer_offer_pct_qpa":     to_float,
    "prevailing_offer_pct_qpa": to_float,
    "provider_offer_pct_median":   to_float,
    "issuer_offer_pct_median":     to_float,
    "prevailing_offer_pct_median": to_float,
    "qpa_pct_median":           to_float,
}


def make_dispute_rows(df: pd.DataFrame, resolver: dict[str, str | None],
                      quarter: str, source_tab: str):
    for record in df.to_dict(orient="records"):
        row: dict = {"quarter": quarter, "source_tab": source_tab}
        for db_col in TAB_DISPUTES_ALIASES:
            src = resolver.get(db_col)
            raw = record.get(src) if src else None
            if db_col in NUMERIC_DISPUTE_FIELDS:
                row[db_col] = NUMERIC_DISPUTE_FIELDS[db_col](raw)
            else:
                row[db_col] = to_str(raw)
        row["health_plan_name_normalized"] = normalize_insurer(row["health_plan_name"])
        # default missing dli_number to dispute_number+row index handled at insert time
        yield tuple(row.get(c) for c in DISPUTES_COLUMNS)


NUMERIC_OFFER_FIELDS = {
    "qpa":               to_float,
    "provider_offer":    to_float,
    "issuer_offer":      to_float,
    "prevailing_offer":  to_float,
}


def make_offer_rows(df: pd.DataFrame, resolver: dict[str, str | None], quarter: str):
    for record in df.to_dict(orient="records"):
        row: dict = {"quarter": quarter}
        for db_col in TAB_OFFERS_ALIASES:
            src = resolver.get(db_col)
            raw = record.get(src) if src else None
            if db_col in NUMERIC_OFFER_FIELDS:
                row[db_col] = NUMERIC_OFFER_FIELDS[db_col](raw)
            else:
                row[db_col] = to_str(raw)
        qpa = row.get("qpa")
        prev = row.get("prevailing_offer")
        pct = None
        if qpa not in (None, 0) and prev is not None:
            pct = (prev / qpa) * 100.0
        row["prevailing_pct_qpa"] = pct
        suspect = 0
        if qpa is not None and qpa < 1.0:
            suspect = 1
        elif pct is not None and pct > 1000:
            suspect = 1
        row["qpa_suspect"] = suspect
        yield tuple(row.get(c) for c in OFFERS_COLUMNS)


# ---------------------------------------------------------------------------
# Insert helpers
# ---------------------------------------------------------------------------

def _placeholders(cols: list[str]) -> str:
    return ",".join("?" for _ in cols)


DISPUTE_INSERT_SQL = (
    f"INSERT OR REPLACE INTO idr_disputes ({','.join(DISPUTES_COLUMNS)}) "
    f"VALUES ({_placeholders(DISPUTES_COLUMNS)})"
)

OFFER_INSERT_SQL = (
    f"INSERT INTO idr_offers ({','.join(OFFERS_COLUMNS)}) "
    f"VALUES ({_placeholders(OFFERS_COLUMNS)})"
)


def _ensure_dli(rows):
    """Disputes are PK'd on (dli_number, quarter). If a row has no dli_number,
    synthesize one so the row still loads. Use dispute_number plus a counter."""
    seen: dict[tuple[str, str], int] = {}
    out = []
    for r in rows:
        r = list(r)
        # column index lookup
        idx_dli = DISPUTES_COLUMNS.index("dli_number")
        idx_disp = DISPUTES_COLUMNS.index("dispute_number")
        idx_quarter = DISPUTES_COLUMNS.index("quarter")
        if not r[idx_dli]:
            base = r[idx_disp] or "anon"
            key = (base, r[idx_quarter])
            seen[key] = seen.get(key, 0) + 1
            r[idx_dli] = f"{base}#{seen[key]}"
        out.append(tuple(r))
    return out


def insert_disputes(conn: sqlite3.Connection, rows: list[tuple]) -> int:
    if not rows:
        return 0
    rows = _ensure_dli(rows)
    conn.executemany(DISPUTE_INSERT_SQL, rows)
    return len(rows)


def insert_offers(conn: sqlite3.Connection, rows: list[tuple]) -> int:
    if not rows:
        return 0
    conn.executemany(OFFER_INSERT_SQL, rows)
    return len(rows)


# ---------------------------------------------------------------------------
# Sheet/file readers
# ---------------------------------------------------------------------------

def read_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet_name, dtype=object, engine="openpyxl")


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=object, low_memory=False)


def expand_zip_to_temp(zip_path: Path, tmpdir: Path) -> list[Path]:
    """Extract a zip and return the files we care about (xlsx/csv)."""
    out = []
    with zipfile.ZipFile(zip_path, "r") as zf:
        for name in zf.namelist():
            lower = name.lower()
            if lower.endswith(".xlsx") or lower.endswith(".csv"):
                target = tmpdir / Path(name).name
                with zf.open(name) as src, open(target, "wb") as dst:
                    while chunk := src.read(1 << 20):
                        dst.write(chunk)
                out.append(target)
    return out


def discover_inputs(puf_dir: Path) -> list[Path]:
    """Return XLSX and ZIP files (skipping the data dictionary)."""
    files: list[Path] = []
    for p in sorted(puf_dir.iterdir()):
        if p.is_dir():
            continue
        if p.name.startswith("~$"):
            continue
        n = p.name.lower()
        if "data-dictionary" in n or "data_dictionary" in n:
            continue
        if n.endswith(".xlsx") or n.endswith(".zip"):
            files.append(p)
    return files


# ---------------------------------------------------------------------------
# Per-file ingest
# ---------------------------------------------------------------------------

def ingest_xlsx(conn: sqlite3.Connection, path: Path, quarter: str,
                file_key: str) -> dict:
    import openpyxl
    print(f"  reading sheets from {path.name}")
    wb = openpyxl.load_workbook(path, read_only=True)
    sheet_names = list(wb.sheetnames)
    wb.close()

    counts = {"oon": 0, "air_ambulance": 0, "offers": 0}
    for sn in sheet_names:
        kind = classify_sheet(sn)
        if kind is None:
            continue
        df = read_sheet(path, sn)
        if df.empty:
            continue
        if kind in ("oon", "air_ambulance"):
            resolver = build_column_resolver(df.columns, TAB_DISPUTES_ALIASES)
            rows = list(make_dispute_rows(df, resolver, quarter, kind))
            inserted = insert_disputes(conn, rows)
        else:
            resolver = build_column_resolver(df.columns, TAB_OFFERS_ALIASES)
            rows = list(make_offer_rows(df, resolver, quarter))
            inserted = insert_offers(conn, rows)
        counts[kind] += inserted
        print(f"    [{quarter}] {sn!r}: {inserted:,} rows")
    return counts


def ingest_csv(conn: sqlite3.Connection, path: Path, quarter: str) -> dict:
    """For zip-extracted CSVs, classify by filename keywords."""
    name = path.name.lower()
    if "air ambulance" in name or "air_ambulance" in name:
        kind = "air_ambulance"
    elif "qpa" in name or "offers" in name:
        kind = "offers"
    elif "emergency" in name or "oon" in name:
        kind = "oon"
    else:
        print(f"    skip (unclassified): {path.name}")
        return {"oon": 0, "air_ambulance": 0, "offers": 0}

    counts = {"oon": 0, "air_ambulance": 0, "offers": 0}
    inserted_total = 0
    # stream large csvs in chunks
    for chunk in pd.read_csv(path, dtype=object, low_memory=False,
                             chunksize=50_000):
        if kind in ("oon", "air_ambulance"):
            resolver = build_column_resolver(chunk.columns, TAB_DISPUTES_ALIASES)
            rows = list(make_dispute_rows(chunk, resolver, quarter, kind))
            inserted_total += insert_disputes(conn, rows)
        else:
            resolver = build_column_resolver(chunk.columns, TAB_OFFERS_ALIASES)
            rows = list(make_offer_rows(chunk, resolver, quarter))
            inserted_total += insert_offers(conn, rows)
    counts[kind] = inserted_total
    print(f"    [{quarter}] {path.name}: {inserted_total:,} rows ({kind})")
    return counts


def ingest_zip(conn: sqlite3.Connection, path: Path, quarter: str,
               file_key: str) -> dict:
    print(f"  expanding {path.name}")
    counts = {"oon": 0, "air_ambulance": 0, "offers": 0}
    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        members = expand_zip_to_temp(path, tmpdir)
        # If a single XLSX is present, prefer it (covers all 3 tabs).
        xlsx = [m for m in members if m.suffix.lower() == ".xlsx"]
        if xlsx:
            for x in xlsx:
                got = ingest_xlsx(conn, x, quarter, file_key)
                for k in counts:
                    counts[k] += got[k]
        else:
            for c in (m for m in members if m.suffix.lower() == ".csv"):
                got = ingest_csv(conn, c, quarter)
                for k in counts:
                    counts[k] += got[k]
    return counts


# ---------------------------------------------------------------------------
# DB helpers
# ---------------------------------------------------------------------------

def ensure_db(conn: sqlite3.Connection) -> None:
    with open(SCHEMA_PATH) as f:
        conn.executescript(f.read())


def already_loaded(conn: sqlite3.Connection, file_key: str) -> bool:
    cur = conn.execute("SELECT 1 FROM ingest_log WHERE file_name = ?", (file_key,))
    return cur.fetchone() is not None


def record_load(conn: sqlite3.Connection, file_key: str, quarter: str,
                counts: dict) -> None:
    conn.execute(
        "INSERT OR REPLACE INTO ingest_log "
        "(file_name, quarter, rows_disputes, rows_air_ambulance, rows_offers, loaded_at) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (file_key, quarter, counts["oon"], counts["air_ambulance"],
         counts["offers"], dt.datetime.utcnow().isoformat(timespec="seconds")),
    )


def refresh_cpt_descriptions(conn: sqlite3.Connection) -> None:
    """Populate cpt_descriptions from the most-frequent description per code."""
    print("  refreshing cpt_descriptions ...")
    conn.execute("DELETE FROM cpt_descriptions")
    conn.execute("""
        INSERT INTO cpt_descriptions (service_code, description, n_seen)
        SELECT service_code, description, n_seen FROM (
            SELECT
                service_code,
                item_description AS description,
                COUNT(*) AS n_seen,
                ROW_NUMBER() OVER (
                    PARTITION BY service_code
                    ORDER BY COUNT(*) DESC
                ) AS rk
            FROM idr_disputes
            WHERE service_code IS NOT NULL AND item_description IS NOT NULL
            GROUP BY service_code, item_description
        )
        WHERE rk = 1
    """)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true",
                    help="reload files even if already in ingest_log")
    ap.add_argument("--file", action="append", default=[],
                    help="ingest only the named file (relative to data/puf/)")
    ap.add_argument("--no-cpt-refresh", action="store_true")
    args = ap.parse_args()

    if not PUF_DIR.exists():
        print(f"PUF directory not found: {PUF_DIR}", file=sys.stderr)
        return 1

    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    ensure_db(conn)

    files = discover_inputs(PUF_DIR)
    if args.file:
        wanted = {Path(f).name for f in args.file}
        files = [f for f in files if f.name in wanted]
    if not files:
        print("No PUF files found.")
        return 0

    print(f"Found {len(files)} input files in {PUF_DIR}")
    grand = {"oon": 0, "air_ambulance": 0, "offers": 0}
    for path in files:
        file_key = path.name
        if already_loaded(conn, file_key) and not args.force:
            print(f"- skip (already loaded): {file_key}")
            continue
        quarter = infer_quarter(file_key)
        print(f"- ingest {file_key}  ->  quarter={quarter}")
        try:
            if path.suffix.lower() == ".zip":
                counts = ingest_zip(conn, path, quarter, file_key)
            else:
                counts = ingest_xlsx(conn, path, quarter, file_key)
            record_load(conn, file_key, quarter, counts)
            conn.commit()
            for k in grand:
                grand[k] += counts[k]
        except Exception as e:
            conn.rollback()
            print(f"  FAILED: {e}", file=sys.stderr)
            raise

    if not args.no_cpt_refresh:
        refresh_cpt_descriptions(conn)
        conn.commit()

    print("\nIngest complete.")
    print(f"  oon disputes     : {grand['oon']:,}")
    print(f"  air ambulance    : {grand['air_ambulance']:,}")
    print(f"  qpa offer rows   : {grand['offers']:,}")
    conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
